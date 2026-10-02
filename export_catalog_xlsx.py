"""Выгрузка каталога товаров в Excel: статистика с диаграммами, общий лист и листы по магазинам.

Запуск: .venv/bin/python export_catalog_xlsx.py [--output catalog_export.xlsx]
База открывается только на чтение — в файл попадает лишь каталог, без пользователей и оплат.
"""

from __future__ import annotations

import argparse
import sqlite3
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.config import get_settings
from database.models import db
from services.product_service import ProductService

ROOT = Path(__file__).resolve().parent

STORES = {
    "mvideo": ("М.Видео", "МВ"),
    "chitaigorod": ("Читай-город", "ЧГ"),
    "detmir": ("Детский мир", "ДМ"),
    "goldapple": ("Золотое яблоко", "ЗЯ"),
}
PRICE_BUCKETS = [
    ("до 1 тыс.", 0, 1000),
    ("1–3 тыс.", 1000, 3000),
    ("3–5 тыс.", 3000, 5000),
    ("5–10 тыс.", 5000, 10000),
    ("10–20 тыс.", 10000, 20000),
    ("20–50 тыс.", 20000, 50000),
    ("50–100 тыс.", 50000, 100000),
    ("от 100 тыс.", 100000, None),
]

# Тестовый набор типичных идей GPT: идея считается покрытой, если поиск бота находит товар в коридоре ±10% от бюджета.
COVERAGE_IDEAS = [
    ("Беспроводные наушники", "беспроводные наушники", 5000, "Электроника"),
    ("AirPods Pro", "apple airpods pro", 25000, "Электроника"),
    ("Портативная колонка JBL", "портативная колонка jbl", 8000, "Электроника"),
    ("Умная колонка с Алисой", "умная колонка яндекс станция", 10000, "Электроника"),
    ("Смарт-часы", "умные часы", 15000, "Электроника"),
    ("Фитнес-браслет", "фитнес браслет", 4000, "Электроника"),
    ("Электронная книга", "электронная книга pocketbook", 15000, "Электроника"),
    ("Игровая консоль PS5", "sony playstation 5", 70000, "Электроника"),
    ("Игра для PS5", "игра ps5", 5000, "Электроника"),
    ("Геймпад", "геймпад dualsense", 7000, "Электроника"),
    ("Игровая мышь", "игровая мышь", 4000, "Электроника"),
    ("Механическая клавиатура", "механическая клавиатура", 7000, "Электроника"),
    ("Планшет", "планшет samsung galaxy tab", 30000, "Электроника"),
    ("Ноутбук для учёбы", "ноутбук для учебы", 50000, "Электроника"),
    ("Смартфон", "смартфон xiaomi", 25000, "Электроника"),
    ("Экшн-камера", "экшн камера", 20000, "Электроника"),
    ("Фотоаппарат моментальной печати", "instax mini", 10000, "Электроника"),
    ("Power bank", "внешний аккумулятор", 3000, "Электроника"),
    ("Беспроводная зарядка", "беспроводная зарядка", 3000, "Электроника"),
    ("Квадрокоптер", "квадрокоптер dji", 40000, "Электроника"),
    ("Проектор", "проектор", 20000, "Электроника"),
    ("Монитор игровой", "игровой монитор", 25000, "Электроника"),
    ("Телевизор", "телевизор 55", 50000, "Электроника"),
    ("Кофемашина", "кофемашина", 30000, "Бытовая техника"),
    ("Капсульная кофеварка", "капсульная кофеварка", 10000, "Бытовая техника"),
    ("Фен Dyson", "фен dyson", 45000, "Бытовая техника"),
    ("Стайлер для волос", "стайлер для волос", 8000, "Бытовая техника"),
    ("Электрическая зубная щётка", "электрическая зубная щетка oral-b", 6000, "Бытовая техника"),
    ("Массажёр", "массажер для шеи", 5000, "Бытовая техника"),
    ("Робот-пылесос", "робот пылесос", 25000, "Бытовая техника"),
    ("Аэрогриль", "аэрогриль", 8000, "Бытовая техника"),
    ("Блендер", "блендер", 6000, "Бытовая техника"),
    ("Электробритва", "электробритва philips", 8000, "Бытовая техника"),
    ("Триммер для бороды", "триммер для бороды", 4000, "Бытовая техника"),
    ("Увлажнитель воздуха", "увлажнитель воздуха", 5000, "Бытовая техника"),
    ("Настольная игра", "настольная игра", 3000, "Хобби и игры"),
    ("Шахматы", "шахматы", 4000, "Хобби и игры"),
    ("Пазл", "пазл", 1500, "Хобби и игры"),
    ("Конструктор LEGO", "конструктор lego", 7000, "Хобби и игры"),
    ("Набор для рисования", "набор для рисования", 3000, "Хобби и игры"),
    ("Телескоп", "телескоп", 15000, "Хобби и игры"),
    ("Парфюм", "парфюмерная вода", 7000, "Красота"),
    ("Набор уходовой косметики", "набор косметики", 4000, "Красота"),
    ("Подарочный набор для мужчин", "подарочный набор для мужчин", 3000, "Красота"),
    ("Книга", "книга", 1500, "Книги"),
    ("Термокружка", "термокружка", 2000, "Дом и аксессуары"),
    ("Плед", "плед", 3000, "Дом и аксессуары"),
    ("Ночник-проектор звёздного неба", "ночник проектор звездное небо", 2500, "Дом и аксессуары"),
    ("Рюкзак", "рюкзак", 5000, "Дом и аксессуары"),
    ("Кожаный кошелёк", "кожаный кошелек", 4000, "Дом и аксессуары"),
    ("Наручные часы", "наручные часы", 10000, "Дом и аксессуары"),
    ("Серьги из серебра", "серьги серебро", 5000, "Украшения"),
    ("Набор инструментов", "набор инструментов", 5000, "Инструменты и дача"),
    ("Шуруповёрт", "шуруповерт", 8000, "Инструменты и дача"),
    ("Набор для гриля", "набор для гриля", 4000, "Инструменты и дача"),
    ("Палатка", "палатка туристическая", 10000, "Инструменты и дача"),
    ("Велосипед", "велосипед", 30000, "Спорт"),
    ("Гантели", "гантели", 4000, "Спорт"),
    ("Коврик для йоги", "коврик для йоги", 2000, "Спорт"),
    ("Электросамокат", "электросамокат", 30000, "Спорт"),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F6BED")
TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=12)
MUTED_FONT = Font(italic=True, color="808080")
LINK_FONT = Font(color="0563C1", underline="single")
PRICE_FORMAT = '#,##0 "₽"'
COUNT_FORMAT = "#,##0"
PRODUCT_COLUMNS = [
    ("Магазин", 16),
    ("Категория", 30),
    ("Бренд", 20),
    ("Название", 70),
    ("Цена", 12),
    ("Ссылка на товар", 45),
    ("Картинка", 30),
    ("Обновлено", 17),
]


def _db_path() -> Path:
    path = Path(get_settings().sqlite_path)
    return path if path.is_absolute() else ROOT / path


def _load_products(path: Path) -> list[dict]:
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    rows = connection.execute(
        "select marketplace, category, store_title, title, price, tracking_link, image_url, updated_at "
        "from affiliateproduct order by marketplace, category, price"
    ).fetchall()
    connection.close()
    return [dict(row) for row in rows]


def _bucket(price: int) -> str:
    for label, low, high in PRICE_BUCKETS:
        if price >= low and (high is None or price < high):
            return label
    return PRICE_BUCKETS[-1][0]


def _store_name(marketplace: str) -> str:
    return STORES.get(marketplace, (marketplace, marketplace))[0]


def _category_name(category: str | None) -> str:
    # Золотое яблоко отдает категории со строчной буквы: «женские ароматы».
    category = (category or "Без категории").strip()
    return category[:1].upper() + category[1:]


def _coverage(path: Path) -> list[dict]:
    # Только чтение: init_db создал бы таблицы и прогнал миграции.
    db.init(f"file:{path}?mode=ro", uri=True)
    db.connect(reuse_if_open=True)
    service = ProductService()
    result = []
    for name, query, budget, kind in COVERAGE_IDEAS:
        offers = service.find_offers([query], min_price=int(budget * 0.9), max_price=int(budget * 1.1))
        result.append(
            {
                "idea": name,
                "type": kind,
                "budget": budget,
                "found": bool(offers),
                "example": offers[0]["label"] if offers else "",
            }
        )
    db.close()
    return result


def _header(ws: Worksheet, row: int, titles: list[str], column: int = 1) -> None:
    for offset, title in enumerate(titles):
        cell = ws.cell(row=row, column=column + offset, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fill_products(ws: Worksheet, products: list[dict]) -> None:
    _header(ws, 1, [title for title, _ in PRODUCT_COLUMNS])
    for index, (_, width) in enumerate(PRODUCT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row, product in enumerate(products, start=2):
        updated = str(product["updated_at"] or "")[:16]
        values = [
            _store_name(product["marketplace"]),
            _category_name(product["category"]),
            product["store_title"] or "",
            product["title"],
            int(product["price"] or 0),
            product["tracking_link"] or "",
            product["image_url"] or "",
            datetime.fromisoformat(updated) if updated else None,
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=value)
        ws.cell(row=row, column=5).number_format = PRICE_FORMAT
        ws.cell(row=row, column=8).number_format = "dd.mm.yyyy hh:mm"
        # Кликабельна только ссылка на товар: у Excel лимит ~65 тыс. гиперссылок на лист.
        link = ws.cell(row=row, column=6)
        if link.value:
            link.hyperlink = link.value
            link.font = LINK_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PRODUCT_COLUMNS))}{len(products) + 1}"


def _section(ws: Worksheet, row: int, title: str) -> int:
    ws.cell(row=row, column=1, value=title).font = SECTION_FONT
    return row + 1


def _no_axis_deletion(chart) -> None:
    # Без этого новые версии Excel прячут оси у диаграмм openpyxl.
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def _labels(*, value: bool = False, percent: bool = False, position: str | None = None) -> DataLabelList:
    # Все флаги явно: иначе Excel добавляет к числу название ряда и категории и рисует выноски.
    labels = DataLabelList(
        showVal=value,
        showPercent=percent,
        showCatName=False,
        showSerName=False,
        showLegendKey=False,
        showLeaderLines=False,
    )
    if position:
        labels.position = position
    return labels


def _single_color(chart, color: str = "4F6BED") -> None:
    # Один ряд — один цвет: без этого Excel красит каждый столбик по-своему.
    chart.varyColors = False
    for series in chart.series:
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color


def _fill_stats(ws: Worksheet, products: list[dict], coverage: list[dict], db_path: Path) -> None:
    total = len(products)
    by_store: dict[str, list[int]] = defaultdict(list)
    by_category: dict[tuple[str, str], list[int]] = defaultdict(list)
    buckets: dict[str, Counter] = defaultdict(Counter)
    for product in products:
        price = int(product["price"] or 0)
        by_store[product["marketplace"]].append(price)
        by_category[(product["marketplace"], _category_name(product["category"]))].append(price)
        buckets[_bucket(price)][product["marketplace"]] += 1
    gift_range = sum(1 for product in products if 1000 <= int(product["price"] or 0) < 10000)
    found = sum(1 for item in coverage if item["found"])

    ws["A1"] = "Каталог товаров бота: статистика"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Источник: {db_path.name} · выгрузка {datetime.now():%d.%m.%Y %H:%M} · "
        "ручной сбор с mvideo.ru, chitai-gorod.ru, detmir.ru, goldapple.ru"
    )
    ws["A2"].font = MUTED_FONT

    row = _section(ws, 4, "Ключевые цифры")
    kpis = [
        ("Товаров в базе", total, "#,##0"),
        ("Магазинов", len(by_store), "0"),
        ("Категорий", len(by_category), "0"),
        ("Товаров за 1–10 тыс. ₽", gift_range / total if total else 0, "0%"),
        ("Тестовых идей получают кнопки", f"{found} из {len(coverage)}", None),
    ]
    for label, value, number_format in kpis:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = Font(bold=True, size=12)
        if number_format:
            cell.number_format = number_format
        row += 1

    row = _section(ws, row + 1, "Товары по магазинам")
    store_header_row = row
    _header(ws, row, ["Магазин", "Товаров", "Доля", "Мин. цена", "Медиана", "Макс. цена"])
    store_order = sorted(by_store, key=lambda market: -len(by_store[market]))
    for market in store_order:
        row += 1
        prices = by_store[market]
        values = [_store_name(market), len(prices), len(prices) / total, min(prices), int(statistics.median(prices)), max(prices)]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=value)
        ws.cell(row=row, column=2).number_format = COUNT_FORMAT
        ws.cell(row=row, column=3).number_format = "0.0%"
        for column in (4, 5, 6):
            ws.cell(row=row, column=column).number_format = PRICE_FORMAT
    store_last_row = row

    row = _section(ws, row + 2, "Распределение товаров по цене")
    price_header_row = row
    _header(ws, row, ["Цена", "Всего", *[_store_name(market) for market in store_order]])
    for label, _, _ in PRICE_BUCKETS:
        row += 1
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=sum(buckets[label].values())).number_format = COUNT_FORMAT
        for offset, market in enumerate(store_order):
            ws.cell(row=row, column=3 + offset, value=buckets[label][market]).number_format = COUNT_FORMAT
    price_last_row = row

    row = _section(ws, row + 2, "Покрытие типичных подарочных идей (поиск бота, бюджет ±10%)")
    coverage_header_row = row
    _header(ws, row, ["Тип идей", "Есть товары", "Нет товаров", "Без кнопок"])
    by_type: dict[str, list[dict]] = defaultdict(list)
    for item in coverage:
        by_type[item["type"]].append(item)
    for kind, items in by_type.items():
        row += 1
        missing = [item["idea"] for item in items if not item["found"]]
        ws.cell(row=row, column=1, value=kind)
        ws.cell(row=row, column=2, value=len(items) - len(missing))
        ws.cell(row=row, column=3, value=len(missing))
        ws.cell(row=row, column=4, value=", ".join(missing) or "—")
    coverage_last_row = row

    row = _section(ws, row + 2, "Товары по категориям")
    category_header_row = row
    _header(ws, row, ["Категория", "Магазин", "Товаров", "Мин. цена", "Медиана", "Макс. цена", "До 5 тыс. ₽"])
    categories = sorted(by_category.items(), key=lambda pair: -len(pair[1]))
    for (market, category), prices in categories:
        row += 1
        values = [
            f"{category} ({STORES.get(market, (market, market))[1]})",
            _store_name(market),
            len(prices),
            min(prices),
            int(statistics.median(prices)),
            max(prices),
            sum(1 for price in prices if price < 5000),
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=value)
        for column in (3, 7):
            ws.cell(row=row, column=column).number_format = COUNT_FORMAT
        for column in (4, 5, 6):
            ws.cell(row=row, column=column).number_format = PRICE_FORMAT
    category_last_row = row

    for column, width in zip("ABCDEFG", (44, 16, 12, 14, 14, 14, 14)):
        ws.column_dimensions[column].width = width

    pie = PieChart()
    pie.title = "Товары по магазинам"
    pie.add_data(Reference(ws, min_col=2, min_row=store_header_row, max_row=store_last_row), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=store_header_row + 1, max_row=store_last_row))
    pie.dataLabels = _labels(percent=True, position="bestFit")
    pie.legend.position = "r"
    pie.height, pie.width = 8, 13
    ws.add_chart(pie, "I4")

    price_chart = BarChart()
    price_chart.title = "Распределение товаров по цене"
    price_chart.y_axis.title = "Товаров, шт."
    price_chart.add_data(Reference(ws, min_col=2, min_row=price_header_row, max_row=price_last_row), titles_from_data=True)
    price_chart.set_categories(Reference(ws, min_col=1, min_row=price_header_row + 1, max_row=price_last_row))
    price_chart.dataLabels = _labels(value=True, position="outEnd")
    price_chart.legend = None
    price_chart.y_axis.majorGridlines = None
    price_chart.gapWidth = 60
    price_chart.height, price_chart.width = 8, 16
    _single_color(price_chart)
    _no_axis_deletion(price_chart)
    ws.add_chart(price_chart, "I21")

    stacked = BarChart()
    stacked.title = "Цены по магазинам"
    stacked.grouping = "stacked"
    stacked.overlap = 100
    stacked.y_axis.title = "Товаров, шт."
    stacked.add_data(
        Reference(ws, min_col=3, max_col=2 + len(store_order), min_row=price_header_row, max_row=price_last_row),
        titles_from_data=True,
    )
    stacked.set_categories(Reference(ws, min_col=1, min_row=price_header_row + 1, max_row=price_last_row))
    stacked.legend.position = "b"
    stacked.gapWidth = 60
    stacked.height, stacked.width = 8, 16
    _no_axis_deletion(stacked)
    ws.add_chart(stacked, "I38")

    coverage_chart = BarChart()
    coverage_chart.type = "bar"
    coverage_chart.grouping = "stacked"
    coverage_chart.overlap = 100
    coverage_chart.title = f"Покрытие тестовых идей: {found} из {len(coverage)}"
    coverage_chart.add_data(
        Reference(ws, min_col=2, max_col=3, min_row=coverage_header_row, max_row=coverage_last_row),
        titles_from_data=True,
    )
    coverage_chart.set_categories(Reference(ws, min_col=1, min_row=coverage_header_row + 1, max_row=coverage_last_row))
    coverage_chart.series[0].graphicalProperties.solidFill = "4CAF50"
    coverage_chart.series[1].graphicalProperties.solidFill = "E57373"
    coverage_chart.x_axis.scaling.orientation = "maxMin"
    # При обратном порядке категорий ось значений уезжает наверх, под заголовок — возвращаем вниз.
    coverage_chart.y_axis.crosses = "max"
    coverage_chart.y_axis.majorGridlines = None
    coverage_chart.legend.position = "b"
    coverage_chart.gapWidth = 50
    coverage_chart.height, coverage_chart.width = 8, 16
    _no_axis_deletion(coverage_chart)
    ws.add_chart(coverage_chart, "I55")

    category_chart = BarChart()
    category_chart.type = "bar"
    category_chart.title = "Количество товаров по категориям"
    category_chart.add_data(
        Reference(ws, min_col=3, min_row=category_header_row, max_row=category_last_row), titles_from_data=True
    )
    category_chart.set_categories(Reference(ws, min_col=1, min_row=category_header_row + 1, max_row=category_last_row))
    category_chart.x_axis.scaling.orientation = "maxMin"
    category_chart.dataLabels = _labels(value=True, position="outEnd")
    category_chart.legend = None
    category_chart.gapWidth = 40
    category_chart.height = max(12, 0.55 * len(categories))
    category_chart.width = 20
    _single_color(category_chart)
    _no_axis_deletion(category_chart)
    # Числа уже подписаны у столбиков: ось значений и сетка только мешают и наезжают на заголовок.
    category_chart.y_axis.delete = True
    category_chart.y_axis.majorGridlines = None
    ws.add_chart(category_chart, "S4")


def _fill_coverage(ws: Worksheet, coverage: list[dict]) -> None:
    _header(ws, 1, ["Идея", "Тип", "Бюджет", "Есть товар", "Пример найденного товара"])
    for row, item in enumerate(coverage, start=2):
        ws.cell(row=row, column=1, value=item["idea"])
        ws.cell(row=row, column=2, value=item["type"])
        ws.cell(row=row, column=3, value=item["budget"]).number_format = PRICE_FORMAT
        status = ws.cell(row=row, column=4, value="да" if item["found"] else "нет")
        status.font = Font(color="2E7D32" if item["found"] else "C62828", bold=True)
        ws.cell(row=row, column=5, value=item["example"])
    for column, width in zip("ABCDE", (34, 20, 12, 12, 50)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(coverage) + 1}"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--output", default=str(ROOT / "catalog_export.xlsx"))
    args = parser.parse_args()

    db_path = _db_path()
    products = _load_products(db_path)
    coverage = _coverage(db_path)

    workbook = Workbook()
    stats = workbook.active
    stats.title = "Статистика"
    _fill_stats(stats, products, coverage, db_path)
    _fill_products(workbook.create_sheet("Все товары"), products)
    by_store: dict[str, list[dict]] = defaultdict(list)
    for product in products:
        by_store[product["marketplace"]].append(product)
    for market in sorted(by_store, key=lambda market: -len(by_store[market])):
        _fill_products(workbook.create_sheet(_store_name(market)), by_store[market])
    _fill_coverage(workbook.create_sheet("Покрытие идей"), coverage)

    output = Path(args.output)
    workbook.save(output)
    print(f"Сохранено: {output} · товаров: {len(products)} · магазинов: {len(by_store)}")


if __name__ == "__main__":
    main()
